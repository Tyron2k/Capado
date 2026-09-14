"""Export and import routines for work package templates.

Provides Excel and CSV exports of templates and their skill requirements and
an importer that creates or updates templates (matched by name) together with
their optional requirements.
"""

import asyncio
import csv
import io

from openpyxl import Workbook
from sqlalchemy import func
from sqlalchemy.ext.asyncio import AsyncSession
from sqlmodel import select

from app.models.skill import Skill, SkillAttribute

from .common import ImportResult, _parse_header


async def _load_templates_export_rows(session: AsyncSession) -> list[tuple]:
    """Load templates with their requirements for export."""
    from app.models.work_package_template import (
        WorkPackageTemplate,
        WorkPackageTemplateRequirement,
    )

    stmt = (
        select(
            WorkPackageTemplate.name,
            WorkPackageTemplate.description,
            Skill.name.label("skill_name"),
            SkillAttribute.name.label("attribute_name"),
            WorkPackageTemplateRequirement.quantity,
        )
        .outerjoin(
            WorkPackageTemplateRequirement,
            WorkPackageTemplateRequirement.template_id == WorkPackageTemplate.id,
        )
        .outerjoin(Skill, Skill.id == WorkPackageTemplateRequirement.skill_id)
        .outerjoin(
            SkillAttribute,
            SkillAttribute.id == WorkPackageTemplateRequirement.skill_attribute_id,
        )
        .order_by(WorkPackageTemplate.name, Skill.name)
    )
    result = await session.execute(stmt)
    return list(result.all())


async def export_templates_xlsx(session: AsyncSession) -> bytes:
    """Export templates and requirements as Excel."""
    rows = await _load_templates_export_rows(session)
    return await asyncio.to_thread(_build_templates_xlsx, rows)


def _build_templates_xlsx(rows: list[tuple]) -> bytes:
    """Build a polished templates Excel workbook grouped by template.

    Visual design: dark header, template group rows, requirement data
    with zebra striping.
    """
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Templates"

    headers = ["Template", "Description", "Skill", "Attribute", "Quantity"]
    total_cols = len(headers)

    # --- Styles ---
    font_header = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    font_template = Font(name="Calibri", bold=True, color="1F3864", size=10)
    font_data = Font(name="Calibri", size=10)

    fill_header = PatternFill(
        start_color="2F5496", end_color="2F5496", fill_type="solid"
    )
    fill_template = PatternFill(
        start_color="D6E4F0", end_color="D6E4F0", fill_type="solid"
    )
    fill_stripe = PatternFill(
        start_color="F7F7F7", end_color="F7F7F7", fill_type="solid"
    )

    thin_side = Side(style="thin", color="D0D0D0")
    border_all = Border(
        left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # --- Header row ---
    ws.row_dimensions[1].height = 22
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_all

    # --- Data rows grouped by template ---
    current_row = 2
    current_template: str | None = None
    row_in_tpl = 0

    for tpl_name, description, skill_name, attr_name, quantity in rows:
        # Template group header
        if tpl_name != current_template:
            current_template = tpl_name
            row_in_tpl = 0
            ws.row_dimensions[current_row].height = 20
            ws.cell(row=current_row, column=1, value=tpl_name)
            ws.cell(row=current_row, column=2, value=description or "")
            for col in range(1, total_cols + 1):
                cell = ws.cell(row=current_row, column=col)
                cell.fill = fill_template
                cell.font = font_template
                cell.border = border_all
                cell.alignment = align_left
            current_row += 1

        # Requirement data row (only if there's a skill)
        if skill_name:
            is_stripe = row_in_tpl % 2 == 1
            ws.cell(row=current_row, column=3, value=skill_name or "")
            ws.cell(row=current_row, column=4, value=attr_name or "")
            ws.cell(row=current_row, column=5, value=quantity if quantity else "")
            for col in range(1, total_cols + 1):
                cell = ws.cell(row=current_row, column=col)
                cell.font = font_data
                cell.border = border_all
                cell.alignment = align_left
                if is_stripe:
                    cell.fill = fill_stripe
            ws.cell(row=current_row, column=5).alignment = align_center
            current_row += 1
            row_in_tpl += 1

    # --- Column widths ---
    col_widths = [24, 30, 18, 18, 10]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Auto-filter & freeze
    if current_row > 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(total_cols)}{current_row - 1}"
    ws.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


async def export_templates_csv(session: AsyncSession) -> str:
    """Export templates and requirements as CSV."""
    rows = await _load_templates_export_rows(session)
    return await asyncio.to_thread(_build_templates_csv, rows)


def _build_templates_csv(rows: list[tuple]) -> str:
    """Build templates CSV content."""
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["Template", "Description", "Skill", "Attribute", "Quantity"])

    for tpl_name, description, skill_name, attr_name, quantity in rows:
        writer.writerow(
            [
                tpl_name,
                description or "",
                skill_name or "",
                attr_name or "",
                quantity if quantity else "",
            ]
        )

    return output.getvalue()


async def import_templates(session: AsyncSession, rows: list[tuple]) -> ImportResult:
    """Import templates and requirements. Matches templates by name.

    Args:
        session: Database session.
        rows: Parsed rows including header.

    Returns:
        Import result with counts.
    """
    from app.models.work_package_template import (
        WorkPackageTemplate,
        WorkPackageTemplateRequirement,
    )

    result = ImportResult()

    if not rows:
        result.errors.append("File is empty.")
        return result

    if len(_parse_header(rows[0])) < 1:
        result.errors.append("Header must have at least: Template.")
        return result

    for row_idx, row in enumerate(rows[1:], start=2):
        cells = [str(cell).strip() if cell else "" for cell in row]
        tpl_name = cells[0] if len(cells) > 0 else ""
        description = cells[1] if len(cells) > 1 else ""
        skill_name = cells[2] if len(cells) > 2 else ""
        attr_name = cells[3] if len(cells) > 3 else ""
        quantity_str = cells[4] if len(cells) > 4 else "1"

        if not tpl_name:
            continue

        # Find or create template
        stmt = select(WorkPackageTemplate).where(
            func.lower(WorkPackageTemplate.name) == tpl_name.lower()
        )
        db_result = await session.execute(stmt)
        template = db_result.scalars().first()

        if template is None:
            template = WorkPackageTemplate(
                name=tpl_name.strip(),
                description=description.strip() or None,
            )
            session.add(template)
            await session.flush()
            result.created += 1
        else:
            if description:
                template.description = description.strip()
                session.add(template)
            result.updated += 1

        # Add requirement if skill is specified
        if skill_name:
            skill_stmt = select(Skill).where(
                func.lower(Skill.name) == skill_name.lower()
            )
            skill_result = await session.execute(skill_stmt)
            skill = skill_result.scalars().first()

            if skill is None:
                result.errors.append(f"Row {row_idx}: Skill '{skill_name}' not found.")
                continue

            skill_attribute_id = None
            if attr_name:
                attr_stmt2 = select(SkillAttribute).where(
                    SkillAttribute.skill_id == skill.id,
                    func.lower(SkillAttribute.name) == attr_name.lower(),
                )
                attr_result2 = await session.execute(attr_stmt2)
                attr = attr_result2.scalars().first()
                if attr is None:
                    result.errors.append(
                        f"Row {row_idx}: Attribute '{attr_name}' not found "
                        f"for skill '{skill_name}'."
                    )
                    continue
                skill_attribute_id = attr.id

            try:
                quantity = int(quantity_str) if quantity_str else 1
            except ValueError:
                quantity = 1

            # Check if requirement already exists
            req_stmt = select(WorkPackageTemplateRequirement).where(
                WorkPackageTemplateRequirement.template_id == template.id,
                WorkPackageTemplateRequirement.skill_id == skill.id,
                WorkPackageTemplateRequirement.skill_attribute_id == skill_attribute_id,
            )
            req_result = await session.execute(req_stmt)
            existing_req = req_result.scalars().first()

            if existing_req is None:
                req = WorkPackageTemplateRequirement(
                    template_id=template.id,
                    skill_id=skill.id,
                    skill_attribute_id=skill_attribute_id,
                    quantity=quantity,
                )
                session.add(req)

    await session.commit()
    return result
