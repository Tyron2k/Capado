"""Export and import routines for projects, work packages, and requirements.

Provides Excel and CSV exports of the project hierarchy (project → work
package → skill requirement) and an importer that creates or updates
projects, their work packages, and optional work package requirements.
"""

import asyncio
import csv
import io
from uuid import UUID

from openpyxl import Workbook
from sqlalchemy import func
from sqlalchemy.ext.asyncio import AsyncSession
from sqlmodel import select

from app.models.skill import Skill, SkillAttribute

from .common import ImportResult, _parse_header


async def _load_projects_export_rows(session: AsyncSession) -> list[tuple]:
    """Load projects with work packages and their requirements for export.

    Each work package with no requirements produces one row with empty
    Skill/Attribute/Quantity. A work package with N requirements produces
    N rows with project/WP data repeated.
    """
    from app.models.project import Project, WorkPackage
    from app.models.work_package_requirement import WorkPackageRequirement

    stmt = (
        select(
            Project.name,
            Project.start_date,
            Project.end_date,
            WorkPackage.name,
            WorkPackage.start_date,
            WorkPackage.end_date,
            Skill.name.label("skill_name"),
            SkillAttribute.name.label("attr_name"),
            WorkPackageRequirement.quantity,
        )
        .outerjoin(WorkPackage, WorkPackage.project_id == Project.id)
        .outerjoin(
            WorkPackageRequirement,
            WorkPackageRequirement.work_package_id == WorkPackage.id,
        )
        .outerjoin(Skill, Skill.id == WorkPackageRequirement.skill_id)
        .outerjoin(
            SkillAttribute,
            SkillAttribute.id == WorkPackageRequirement.skill_attribute_id,
        )
        .order_by(Project.name, WorkPackage.start_date, Skill.name)
    )
    result = await session.execute(stmt)
    return list(result.all())


async def export_projects_xlsx(session: AsyncSession) -> bytes:
    """Export projects, work packages, and requirements as Excel."""
    rows = await _load_projects_export_rows(session)
    return await asyncio.to_thread(_build_projects_xlsx, rows)


def _build_projects_xlsx(rows: list[tuple]) -> bytes:
    """Build a polished projects Excel workbook grouped by project and work package.

    Visual design matches the skill matrix: dark header, project group rows,
    work package sub-groups, zebra striping, borders, auto-filter, freeze panes.
    """
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Projects"

    headers = [
        "Project",
        "Project Start",
        "Project End",
        "Work Package",
        "WP Start",
        "WP End",
        "Skill",
        "Attribute",
        "Quantity",
    ]
    total_cols = len(headers)

    # --- Styles ---
    font_header = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    font_project = Font(name="Calibri", bold=True, color="1F3864", size=10)
    font_wp = Font(name="Calibri", bold=True, color="4A4A4A", size=10)
    font_data = Font(name="Calibri", size=10)

    fill_header = PatternFill(
        start_color="2F5496", end_color="2F5496", fill_type="solid"
    )
    fill_project = PatternFill(
        start_color="D6E4F0", end_color="D6E4F0", fill_type="solid"
    )
    fill_wp = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid")
    fill_stripe = PatternFill(
        start_color="F7F7F7", end_color="F7F7F7", fill_type="solid"
    )

    thin_side = Side(style="thin", color="D0D0D0")
    border_all = Border(
        left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # --- Header row ---
    ws.row_dimensions[1].height = 22
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_all

    # --- Data rows grouped by project → work package ---
    current_row = 2
    current_project: str | None = None
    current_wp: str | None = None
    row_in_wp = 0

    for (
        project_name,
        p_start,
        p_end,
        wp_name,
        wp_start,
        wp_end,
        skill_name,
        attr_name,
        quantity,
    ) in rows:
        # Project group header
        if project_name != current_project:
            current_project = project_name
            current_wp = None
            ws.row_dimensions[current_row].height = 20
            ws.cell(row=current_row, column=1, value=project_name)
            ws.cell(
                row=current_row,
                column=2,
                value=p_start.isoformat() if p_start else "",
            )
            ws.cell(
                row=current_row,
                column=3,
                value=p_end.isoformat() if p_end else "",
            )
            for col in range(1, total_cols + 1):
                cell = ws.cell(row=current_row, column=col)
                cell.fill = fill_project
                cell.font = font_project
                cell.border = border_all
                cell.alignment = align_left
            current_row += 1

        # Work package sub-header
        if wp_name and wp_name != current_wp:
            current_wp = wp_name
            row_in_wp = 0
            ws.cell(row=current_row, column=4, value=wp_name)
            ws.cell(
                row=current_row,
                column=5,
                value=wp_start.isoformat() if wp_start else "",
            )
            ws.cell(
                row=current_row,
                column=6,
                value=wp_end.isoformat() if wp_end else "",
            )
            for col in range(1, total_cols + 1):
                cell = ws.cell(row=current_row, column=col)
                cell.fill = fill_wp
                cell.font = font_wp
                cell.border = border_all
                cell.alignment = align_left
            current_row += 1

        # Requirement data row (only if there's a skill)
        if skill_name:
            is_stripe = row_in_wp % 2 == 1
            ws.cell(row=current_row, column=7, value=skill_name or "")
            ws.cell(row=current_row, column=8, value=attr_name or "")
            ws.cell(row=current_row, column=9, value=quantity if quantity else "")
            for col in range(1, total_cols + 1):
                cell = ws.cell(row=current_row, column=col)
                cell.font = font_data
                cell.border = border_all
                cell.alignment = align_left
                if is_stripe:
                    cell.fill = fill_stripe
            ws.cell(row=current_row, column=9).alignment = align_center
            current_row += 1
            row_in_wp += 1

    # --- Column widths ---
    col_widths = [22, 13, 13, 26, 13, 13, 18, 18, 10]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Auto-filter & freeze
    if current_row > 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(total_cols)}{current_row - 1}"
    ws.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


async def export_projects_csv(session: AsyncSession) -> str:
    """Export projects, work packages, and requirements as CSV."""
    rows = await _load_projects_export_rows(session)
    return await asyncio.to_thread(_build_projects_csv, rows)


def _build_projects_csv(rows: list[tuple]) -> str:
    """Build projects CSV content (CPU-bound, runs in thread pool)."""
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(
        [
            "Project",
            "Project Start",
            "Project End",
            "Work Package",
            "WP Start",
            "WP End",
            "Skill",
            "Attribute",
            "Quantity",
        ]
    )

    for (
        project_name,
        p_start,
        p_end,
        wp_name,
        wp_start,
        wp_end,
        skill_name,
        attr_name,
        quantity,
    ) in rows:
        writer.writerow(
            [
                project_name,
                p_start.isoformat() if p_start else "",
                p_end.isoformat() if p_end else "",
                wp_name or "",
                wp_start.isoformat() if wp_start else "",
                wp_end.isoformat() if wp_end else "",
                skill_name or "",
                attr_name or "",
                quantity if quantity else "",
            ]
        )

    return output.getvalue()


async def import_projects(session: AsyncSession, rows: list[tuple]) -> ImportResult:
    """Import projects, work packages, and optional requirements.

    Each row must have at least Project, Project Start, Project End. Optional
    columns: Work Package, WP Start, WP End, Skill, Attribute, Quantity.

    Skills must already exist — if not found, an error is reported for that row.
    Attribute is optional (null means any attribute satisfies the requirement).
    Quantity defaults to 1 if empty.

    Args:
        session: Database session.
        rows: Parsed rows including header.

    Returns:
        Import result with counts.
    """
    from datetime import date as date_type

    from app.models.project import Project, WorkPackage
    from app.models.work_package_requirement import WorkPackageRequirement

    result = ImportResult()

    if not rows:
        result.errors.append("File is empty.")
        return result

    if len(_parse_header(rows[0])) < 3:
        result.errors.append(
            "Header must have at least: Project, Project Start, Project End."
        )
        return result

    # Pre-load skill caches for requirement resolution
    skill_stmt = select(Skill)
    skill_result = await session.execute(skill_stmt)
    skills_by_name: dict[str, Skill] = {
        s.name.lower(): s for s in skill_result.scalars().all()
    }

    attr_stmt = select(SkillAttribute)
    attr_result = await session.execute(attr_stmt)
    attrs_by_key: dict[tuple[UUID, str], SkillAttribute] = {
        (a.skill_id, a.name.lower()): a for a in attr_result.scalars().all()
    }

    # Track existing requirements to skip duplicates
    req_stmt = select(WorkPackageRequirement)
    req_result = await session.execute(req_stmt)
    existing_reqs: set[tuple[UUID, UUID, UUID | None]] = {
        (r.work_package_id, r.skill_id, r.skill_attribute_id)
        for r in req_result.scalars().all()
    }

    for row_idx, row in enumerate(rows[1:], start=2):
        cells = [str(cell).strip() if cell else "" for cell in row]
        project_name = cells[0] if len(cells) > 0 else ""
        p_start_str = cells[1] if len(cells) > 1 else ""
        p_end_str = cells[2] if len(cells) > 2 else ""
        wp_name = cells[3] if len(cells) > 3 else ""
        wp_start_str = cells[4] if len(cells) > 4 else ""
        wp_end_str = cells[5] if len(cells) > 5 else ""
        skill_name = cells[6] if len(cells) > 6 else ""
        attr_name = cells[7] if len(cells) > 7 else ""
        quantity_str = cells[8] if len(cells) > 8 else ""

        if not project_name:
            continue

        # Parse project dates
        try:
            p_start = date_type.fromisoformat(p_start_str) if p_start_str else None
            p_end = date_type.fromisoformat(p_end_str) if p_end_str else None
        except ValueError:
            result.errors.append(
                f"Row {row_idx}: Invalid project date format (use YYYY-MM-DD)."
            )
            continue

        if not p_start or not p_end:
            result.errors.append(
                f"Row {row_idx}: Project Start and Project End are required."
            )
            continue

        # Find or create project
        stmt = select(Project).where(func.lower(Project.name) == project_name.lower())
        db_result = await session.execute(stmt)
        project = db_result.scalars().first()

        if project is None:
            project = Project(
                name=project_name.strip(),
                start_date=p_start,
                end_date=p_end,
            )
            session.add(project)
            await session.flush()
            result.created += 1
        else:
            project.start_date = p_start
            project.end_date = p_end
            session.add(project)
            result.updated += 1

        # Create work package if specified
        wp = None
        if wp_name:
            try:
                wp_start = (
                    date_type.fromisoformat(wp_start_str) if wp_start_str else p_start
                )
                wp_end = date_type.fromisoformat(wp_end_str) if wp_end_str else p_end
            except ValueError:
                result.errors.append(
                    f"Row {row_idx}: Invalid work package date format (use YYYY-MM-DD)."
                )
                continue

            # Check if WP already exists for this project
            wp_stmt = select(WorkPackage).where(
                WorkPackage.project_id == project.id,
                func.lower(WorkPackage.name) == wp_name.lower(),
            )
            wp_result = await session.execute(wp_stmt)
            wp = wp_result.scalars().first()

            if wp is None:
                wp = WorkPackage(
                    project_id=project.id,
                    name=wp_name.strip(),
                    start_date=wp_start,
                    end_date=wp_end,
                )
                session.add(wp)
                await session.flush()
            else:
                wp.start_date = wp_start
                wp.end_date = wp_end
                session.add(wp)

        # Handle optional requirement
        if skill_name and wp:
            skill = skills_by_name.get(skill_name.lower())
            if skill is None:
                result.errors.append(f"Row {row_idx}: Skill '{skill_name}' not found.")
                continue

            skill_attribute_id: UUID | None = None
            if attr_name:
                attr_key = (skill.id, attr_name.lower())
                attribute = attrs_by_key.get(attr_key)
                if attribute is None:
                    result.errors.append(
                        f"Row {row_idx}: Attribute '{attr_name}' not found "
                        f"for skill '{skill_name}'."
                    )
                    continue
                skill_attribute_id = attribute.id

            try:
                quantity = int(quantity_str) if quantity_str else 1
            except ValueError:
                quantity = 1

            # Skip duplicate requirements
            req_key = (wp.id, skill.id, skill_attribute_id)
            if req_key not in existing_reqs:
                session.add(
                    WorkPackageRequirement(
                        work_package_id=wp.id,
                        skill_id=skill.id,
                        skill_attribute_id=skill_attribute_id,
                        quantity=quantity,
                    )
                )
                existing_reqs.add(req_key)

    await session.commit()
    return result
