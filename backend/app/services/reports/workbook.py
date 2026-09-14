"""Excel workbooks meant to be READ, not re-imported.

Deliberately separate from :mod:`app.services.import_export`, which exists for round-trip data
exchange: those files are exported, edited and imported back, so their column names and shapes are
a contract. A management report has no such contract — it is aggregated, formatted for a recipient
who will never re-import it, and free to change. Putting the two in one module would eventually
break somebody's import because a report needed another column.

**The failure that makes an export useless is writing numbers as text.** A recipient's first act is
to sum a column or build a pivot table, and a sheet of numeric strings silently refuses both — no
error, just wrong or empty totals. Every numeric cell here is written as a Python int or float and
every date as a ``date``, so Excel receives real types. That is the single most important thing in
this file and the reason for the ``write_row`` helper rather than ad-hoc ``ws.append`` calls.

Second: a report states when it was generated and what it covers. A spreadsheet that circulates by
email without those two facts gets quoted six weeks later as if it were current.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
_HEADER_FONT = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=14)

# Column widths are set from content length rather than left to Excel's default, because the
# default truncates every name and a recipient who has to widen ten columns before reading
# treats the file as broken.
_MIN_WIDTH = 8
_MAX_WIDTH = 50


@dataclass(frozen=True)
class Column:
    """One report column.

    Attributes:
        header: Column heading.
        number_format: openpyxl number format, e.g. ``"0.0"`` or ``"DD.MM.YYYY"``. Applied to
            the cell rather than baked into a string, which is what keeps the value numeric.
    """

    header: str
    number_format: str | None = None


def write_title(ws: Worksheet, title: str, subtitle: str) -> int:
    """Write the title block and return the next free row.

    The subtitle carries the coverage and generation date. Not optional: a spreadsheet
    circulating by email without them gets quoted six weeks later as if it were current.
    """
    ws.cell(row=1, column=1, value=title).font = _TITLE_FONT
    ws.cell(row=2, column=1, value=subtitle)
    return 4


def write_header(ws: Worksheet, columns: Sequence[Column], row: int) -> int:
    """Write the header row and freeze it. Returns the next free row.

    Frozen because a report long enough to scroll is a report where the reader loses track of
    which column is which, and then misreads a number rather than asking.
    """
    for index, column in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=index, value=column.header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    # Set by A1-style reference rather than by cell object: openpyxl types freeze_panes as
    # str | Cell | None, and handing it the Cell returned by ws.cell() is a type error even
    # though it works at runtime.
    ws.freeze_panes = f"A{row + 1}"
    return row + 1


def write_row(
    ws: Worksheet,
    columns: Sequence[Column],
    values: Sequence[object],
    row: int,
) -> int:
    """Write one data row with the column's number format applied. Returns the next row.

    Values are passed through unchanged — no str() anywhere. An int stays an int, a date stays a
    date, and None leaves the cell genuinely empty rather than holding the text "None".
    """
    for index, (column, value) in enumerate(
        zip(columns, values, strict=False), start=1
    ):
        cell = ws.cell(row=row, column=index, value=value)
        if column.number_format and value is not None:
            cell.number_format = column.number_format
    return row + 1


def autosize(ws: Worksheet, columns: Sequence[Column], header_row: int) -> None:
    """Set column widths from the widest cell in each column.

    Measured over the actual written cells rather than guessed from the header, because the
    header is usually the shortest string in the column.
    """
    for index, column in enumerate(columns, start=1):
        longest = len(column.header)
        for row in range(header_row + 1, ws.max_row + 1):
            value = ws.cell(row=row, column=index).value
            if value is not None:
                longest = max(longest, len(str(value)))
        ws.column_dimensions[get_column_letter(index)].width = min(
            _MAX_WIDTH, max(_MIN_WIDTH, longest + 2)
        )


def new_workbook() -> Workbook:
    """A workbook with the default sheet removed.

    openpyxl creates a sheet named "Sheet" that every caller then renames or leaves behind. An
    empty stray sheet in a report reads as something having gone wrong.
    """
    workbook = Workbook()
    default = workbook.active
    if default is not None:
        workbook.remove(default)
    return workbook


def save(workbook: Workbook) -> bytes:
    """Serialize to bytes for an HTTP response."""
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def generated_subtitle(covers: str, now: datetime | None = None) -> str:
    """The subtitle line: what this covers and when it was made."""
    moment = now or datetime.now()
    return f"{covers} — erstellt am {moment.strftime('%d.%m.%Y %H:%M')}"


def filename_for(report: str, today: date) -> str:
    """A filename that sorts chronologically and says what it is.

    ISO date first so a folder of these sorts by date in every file manager, which a
    ``26.08.2026`` prefix would not.
    """
    return f"{today.isoformat()}-capado-{report}.xlsx"
