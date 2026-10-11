"""Read real XLSX exports back: user text must never become a formula."""

import io
from datetime import date, datetime
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app import models as m
from app.models.project import ProjectFolder
from app.services.import_export.common import (
    _build_flat_resource_xlsx,
    _build_skill_matrix_xlsx,
    _MatrixData,
)
from app.services.import_export.projects import _build_projects_xlsx
from app.services.import_export.templates import _build_templates_xlsx
from app.services.reports.projects import build_project_report
from app.services.reports.utilization import build_utilization_report
from app.services.reports.workbook import (
    Column,
    new_workbook,
    save,
    write_header,
    write_row,
    write_title,
)


def _read(data):
    return load_workbook(io.BytesIO(data), data_only=False)


def _assert_text(cell, text):
    assert cell.value == text
    assert cell.data_type == "s"


@pytest.mark.parametrize(
    "name",
    [
        "Anna Beispiel",
        "=1+1",
        '=HYPERLINK("https://example.test","Test")',
        "#N/A",
        "+1",
        "-1",
        "@Name",
        "Änne 李 & <Team>",
    ],
)
def test_report_writes_text_without_changing_it_or_other_cell_types(name):
    workbook = new_workbook()
    sheet = workbook.create_sheet("Report")
    write_title(sheet, name, name)
    columns = [
        Column(name),
        Column("Date", "DD.MM.YYYY"),
        Column("Hours", "0.0"),
        Column("Flag"),
        Column("Optional"),
    ]
    write_header(sheet, columns, 4)
    write_row(sheet, columns, [name, date(2026, 10, 11), 37.5, True, None], 5)
    result = _read(save(workbook)).active
    for address in ["A1", "A2", "A4", "A5"]:
        _assert_text(result[address], name)
    assert result["B5"].data_type == "d"
    assert result["B5"].value == datetime(2026, 10, 11)
    assert result["B5"].number_format == "DD.MM.YYYY"
    assert result["C5"].data_type == "n" and result["C5"].value == 37.5
    assert result["C5"].number_format == "0.0"
    assert result["D5"].data_type == "b" and result["D5"].value is True
    assert result["E5"].value is None
    assert result.freeze_panes == "A5"
    assert result["A1"].font.bold and result["A4"].font.bold


def test_matrix_names_headers_and_merged_groups_stay_text():
    first, second = uuid4(), uuid4()
    data = _MatrixData(
        columns=[("=Skill", "=Attribute"), ("=Skill", "Unicode 李")],
        resources=[("=Person", "=Group", {first}), ("Anna", "=Group", {second})],
        attr_id_order=[first, second],
    )
    sheet = _read(_build_skill_matrix_xlsx(data, "Personnel")).active
    for address, name in {
        "C1": "=Skill",
        "C2": "=Attribute",
        "A3": "=Group",
        "A4": "=Person",
        "B4": "=Group",
    }.items():
        _assert_text(sheet[address], name)
    assert sheet["C4"].value == "X"
    assert "C1:D1" in sheet.merged_cells and "A3:D3" in sheet.merged_cells
    assert sheet.freeze_panes == "C3"
    assert sheet["C1"].fill.fgColor.rgb == "002F5496"


@pytest.mark.parametrize("sheet_title", ["Personnel", "Infrastructure"])
def test_flat_resource_xlsx_roundtrip_preserves_all_names(sheet_title):
    values = ("=Person", "=Group", "=Skill", "=Attribute", "=Site")
    data = _build_flat_resource_xlsx(
        [values, ("李", "Team", None, None, None)], sheet_title
    )
    sheet = _read(data).active
    for cell, value in zip(sheet[2], values, strict=True):
        _assert_text(cell, value)
    # Imports read cached values. Formula cells would lose their name here.
    imported = load_workbook(io.BytesIO(data), data_only=True).active
    assert tuple(cell.value for cell in imported[2]) == values
    assert tuple(cell.value for cell in imported[3]) == ("李", "Team", None, None, None)


def test_project_and_template_export_names_stay_text():
    day = date(2026, 10, 11)
    project = _read(
        _build_projects_xlsx(
            [("=Project", day, day, "=Package", day, day, "=Skill", "=Attribute", 3)]
        )
    ).active
    template = _read(
        _build_templates_xlsx(
            [("=Template", "=Description", "=Skill", "=Attribute", 3)]
        )
    ).active
    for sheet, names in [
        (project, {"=Project", "=Package", "=Skill", "=Attribute"}),
        (template, {"=Template", "=Description", "=Skill", "=Attribute"}),
    ]:
        seen = set()
        for row in sheet:
            for cell in row:
                assert cell.data_type != "f"
                if cell.value in names:
                    _assert_text(cell, cell.value)
                    seen.add(cell.value)
        assert seen == names
        assert any(
            cell.value == 3 and cell.data_type == "n" for row in sheet for cell in row
        )


async def test_reports_preserve_database_names_and_dates(db_session):
    day = date(2026, 10, 12)
    customer = m.Customer(name="=Customer")
    group = m.ResourceGroup(name="=Group", resource_type="personal")
    profile = m.WorkWeekProfile(name="Standard", is_default=True)
    db_session.add_all([customer, group, profile])
    await db_session.flush()
    folder = ProjectFolder(name="=Folder", customer_id=customer.id)
    person = m.PersonalResource(name="=Person", group_id=group.id)
    db_session.add_all([folder, person])
    await db_session.flush()
    project = m.Project(
        name="=Project", folder_id=folder.id, start_date=day, end_date=day
    )
    db_session.add(project)
    await db_session.commit()

    sheet = _read(await build_project_report(db_session))["Projekte"]
    for address, value in {
        "A5": "=Folder",
        "B5": "=Project",
        "C5": "=Customer",
    }.items():
        _assert_text(sheet[address], value)
    assert sheet["G5"].data_type == "d" and sheet["G5"].number_format == "DD.MM.YYYY"
    utilization = _read(await build_utilization_report(db_session, day, day))
    cells = [cell for sheet in utilization for row in sheet for cell in row]
    for name in ["=Group", "=Person"]:
        matches = [cell for cell in cells if cell.value == name]
        assert matches
        for cell in matches:
            _assert_text(cell, name)
    assert all(cell.data_type != "f" for cell in cells)
    for entity, name in [
        (customer, "=Customer"),
        (group, "=Group"),
        (person, "=Person"),
        (project, "=Project"),
        (folder, "=Folder"),
    ]:
        await db_session.refresh(entity)
        assert entity.name == name
