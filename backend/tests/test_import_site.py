"""Importing the site column: what it means when it is absent, empty, or names an unknown site.

WRITTEN BEFORE THE IMPLEMENTATION, on purpose. Twice in this feature a test passed against broken
code — once because the value under test was already null before the mapping ran. A rule that
REJECTS data is the worst place for that to happen: a toothless test certifies a barrier that is not
there, and the failure only shows up as a wrong plan months later.

THE THREE CASES, and why they are not the same:

- **Column absent** — leave the site alone. Every file exported before this column existed has no
  Site header, and re-importing one must not silently unfile every resource in it.
- **Column present, cell empty** — REMOVE the site. Without this, an operator can add a site through
  the Excel round-trip but never take one away, which is the same gap as using None for "leave
  alone" in the API.
- **Column present, unknown name** — reject the row. This deliberately differs from Group, which the
  importer AUTO-CREATES. A site owns the holiday calendar: "Werk Amendorf" with a missing letter
  would create a second plant with NO holidays, and every resource landing there silently loses the
  calendar that drives its capacity. An untidy group list is cosmetic; a wrong holiday calendar is a
  wrong plan.
"""

from __future__ import annotations

import pytest

from app.services.import_export.shape import ImportShape, detect_import_shape


class TestHeaderDetection:
    """A Site column must not make a valid file unrecognisable."""

    def test_a_file_with_a_site_column_is_still_the_flat_shape(self):
        rows = [
            ("Name", "Group", "Skill", "Attribute", "Site"),
            ("Müller", "Lackierer", "", "", ""),
        ]
        assert detect_import_shape(rows) is ImportShape.FLAT

    def test_a_german_site_header_is_accepted_too(self):
        """Hand-built German files are real: the in-app help documented German column names."""
        rows = [
            ("Name", "Gruppe", "Betriebsstätte"),
            ("Müller", "Lackierer", "Werk Ammendorf"),
        ]
        assert detect_import_shape(rows) is ImportShape.FLAT

    def test_a_file_without_a_site_column_is_unchanged(self):
        rows = [
            ("Name", "Group", "Skill", "Attribute"),
            ("Müller", "Lackierer", "", ""),
        ]
        assert detect_import_shape(rows) is ImportShape.FLAT


class TestSiteColumnIndex:
    """Finding the column by NAME rather than by position, and why."""

    def test_the_site_column_is_found_wherever_it_sits(self):
        """Positional would break every existing file: the importer reads 0..3 by index today, so a
        Site column wedged in at index 2 would be read as a Skill. By name it does not matter."""
        from app.services.import_export.common import _site_column_index

        assert _site_column_index(["Name", "Group", "Skill", "Attribute", "Site"]) == 4
        assert _site_column_index(["Name", "Group", "Betriebsstätte"]) == 2

    def test_no_site_column_reports_none(self):
        from app.services.import_export.common import _site_column_index

        assert _site_column_index(["Name", "Group", "Skill", "Attribute"]) is None

    def test_detection_ignores_case_and_padding(self):
        from app.services.import_export.common import _site_column_index

        assert _site_column_index(["Name", "Group", "  SITE  "]) == 2


class TestResolvingASiteName:
    """The rejection, which is the whole point of this file."""

    def test_a_known_name_resolves_case_insensitively(self):
        from app.services.import_export.common import resolve_site

        known = {"Werk Ammendorf": "id-a", "Werk Leipzig": "id-b"}
        assert resolve_site("Werk Ammendorf", known) == ("id-a", None)
        assert resolve_site("WERK LEIPZIG", known) == ("id-b", None)

    def test_an_empty_name_clears_the_site_without_an_error(self):
        from app.services.import_export.common import resolve_site

        assert resolve_site("", {"Werk Ammendorf": "id-a"}) == (None, None)

    def test_an_unknown_name_is_refused_and_never_created(self):
        from app.services.import_export.common import resolve_site

        site_id, error = resolve_site("Werk Amendorf", {"Werk Ammendorf": "id-a"})
        assert site_id is None
        assert error is not None

    def test_the_error_names_the_input_and_the_known_sites(self):
        """An operator staring at a rejected row needs to see their typo next to the real name."""
        from app.services.import_export.common import resolve_site

        _, error = resolve_site("Werk Amendorf", {"Werk Ammendorf": "id-a"})
        assert error is not None
        assert "Werk Amendorf" in error
        assert "Ammendorf" in error

    def test_an_unknown_name_is_refused_even_when_nothing_is_known(self):
        from app.services.import_export.common import resolve_site

        site_id, error = resolve_site("Werk Ammendorf", {})
        assert site_id is None
        assert error is not None


class TestTheExportSideOfTheRoundTrip:
    """The editable Excel export remains compatible with the legacy importer."""

    def test_the_flat_excel_header_and_site_match_the_importer(self):
        import io

        from openpyxl import load_workbook

        from app.services.import_export.common import (
            _build_flat_resource_xlsx,
            _site_column_index,
        )

        workbook = load_workbook(
            io.BytesIO(
                _build_flat_resource_xlsx(
                    [
                        ("Müller", "Lackierer", None, None, "Werk Ammendorf"),
                        ("Schmidt", "Lackierer", None, None, None),
                    ],
                    "Personnel",
                )
            ),
            read_only=True,
            data_only=True,
        )
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
            assert detect_import_shape(rows) is ImportShape.FLAT
            index = _site_column_index(rows[0])
            assert index is not None
            assert rows[1][index] == "Werk Ammendorf"
            assert rows[2][index] is None
        finally:
            workbook.close()


@pytest.mark.parametrize(
    "header", ["Site", "site", "Betriebsstätte", "betriebsstaette"]
)
def test_every_accepted_spelling_is_found(header: str):
    """Including the ASCII fallback, because a CSV round-trip through a mis-encoded editor loses ä."""
    from app.services.import_export.common import _site_column_index

    assert _site_column_index(["Name", "Group", header]) == 2
